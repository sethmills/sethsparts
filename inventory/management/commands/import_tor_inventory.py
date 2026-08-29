import re

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from inventory.models import Category, Container, Drawer, Location, Part, StockItem

# Container types to import — everything that isn't a literal cardboard moving box.
# Matched case-insensitively against the spreadsheet's 'Box Size' column.
INCLUDED_CONTAINER_TYPES = {
    "custom",
    "large black tote",
    "small black tote",
    "medium black tote",
    "extra large black tote",
    "large black tote with flip lid",
    "large black plastic",
    "extra long black plastic",
    "long green plastic",
    "medium green plastic",
    "black long",
    "small clear tote",
    "small clear plastic",
    "medium clear plastic",
    "medium plastic",
    "cabinets",
}

DRAWER_NOTE_RE = re.compile(r"^drawer\s+([a-z]\d+)$", re.IGNORECASE)

# Very small first-pass keyword heuristic for categorizing parts on import.
# Refine actual categorization in /admin — this just gives Phase 3+ a starting queue.
CATEGORY_KEYWORDS = {
    "Electronics": [
        "adafruit", "feather", "raspi", "pi zero", "arduino", "rp2040", "resistor",
        "capacitor", "diode", "transistor", "relay", "sensor", "led", "wire", "cable",
        "connector", "terminal", "header", "breadboard", "solder", "circuit", "pcb",
        "battery", "motor", "servo", "stepper", "board", "oled", "usb", "jst", "wago",
        "microswitch", "power supply", "ide cable",
    ],
    "3D Printing": ["fillament", "filament", "3d print", "nozzle", "extruder", "print bed"],
    "Tools": ["screwdriver", "wrench", "drill", "saw", "pliers", "hammer", "socket set"],
}

ADAFRUIT_MARKER = "adafruit"


class Command(BaseCommand):
    help = "One-time import of TOR Inventory.xlsx into the inventory app."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str)

    def handle(self, *args, **options):
        path = options["xlsx_path"]
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f"File not found: {path}") from exc

        contents_ws = wb["Box Contents"]
        rows = list(contents_ws.iter_rows(min_row=2, values_only=True))

        dims_by_box_number = self._load_dimensions(wb)

        stats = {
            "rows_total": len(rows),
            "rows_skipped_empty": 0,
            "rows_skipped_excluded_type": 0,
            "containers_created": 0,
            "drawers_created": 0,
            "parts_created": 0,
            "stock_items_created": 0,
        }

        with transaction.atomic():
            containers_by_number = {}
            drawers_by_key = {}
            categories_by_name = {}
            parts_by_normalized_name = {}

            for room, box_number, box_size, item, num_items, notes, extra in rows:
                if box_number is None or item is None:
                    stats["rows_skipped_empty"] += 1
                    continue

                normalized_type = (box_size or "").strip().lower()
                if normalized_type not in INCLUDED_CONTAINER_TYPES:
                    stats["rows_skipped_excluded_type"] += 1
                    continue

                box_number = int(box_number)
                location = self._get_location(room)

                container = containers_by_number.get(box_number)
                if container is None:
                    container, created = Container.objects.get_or_create(
                        number=box_number,
                        defaults={
                            "container_type": (box_size or "").strip(),
                            "location": location,
                            "dimensions": dims_by_box_number.get(box_number, ""),
                            "notes": (extra or "").strip(),
                        },
                    )
                    if created:
                        stats["containers_created"] += 1
                    containers_by_number[box_number] = container

                drawer = None
                source_notes = ""
                if notes:
                    notes_str = str(notes).strip()
                    match = DRAWER_NOTE_RE.match(notes_str)
                    if match:
                        drawer_key = (box_number, notes_str.lower())
                        drawer = drawers_by_key.get(drawer_key)
                        if drawer is None:
                            drawer, created = Drawer.objects.get_or_create(
                                container=container, label=notes_str
                            )
                            if created:
                                stats["drawers_created"] += 1
                            drawers_by_key[drawer_key] = drawer
                    else:
                        source_notes = notes_str

                item_name = str(item).strip()
                normalized_name = item_name.lower()
                part = parts_by_normalized_name.get(normalized_name)
                if part is None:
                    category = self._categorize(item_name, categories_by_name)
                    part, created = Part.objects.get_or_create(
                        normalized_name=normalized_name,
                        defaults={
                            "name": item_name,
                            "category": category,
                            "is_electronic": category is not None and category.name == "Electronics",
                            "enrichment_status": (
                                Part.ENRICHMENT_PENDING
                                if ADAFRUIT_MARKER in normalized_name
                                else Part.ENRICHMENT_NOT_NEEDED
                            ),
                        },
                    )
                    if created:
                        stats["parts_created"] += 1
                    parts_by_normalized_name[normalized_name] = part

                quantity_raw = "" if num_items is None else str(num_items).strip()
                quantity = self._parse_quantity(num_items)

                StockItem.objects.create(
                    part=part,
                    container=container,
                    drawer=drawer,
                    quantity=quantity,
                    quantity_raw=quantity_raw,
                    source_notes=source_notes,
                )
                stats["stock_items_created"] += 1

        self.stdout.write(self.style.SUCCESS("Import complete:"))
        for key, value in stats.items():
            self.stdout.write(f"  {key}: {value}")

    def _get_location(self, room):
        if not room:
            return None
        name = str(room).strip()
        if not name:
            return None
        location, _ = Location.objects.get_or_create(name=name)
        return location

    def _categorize(self, item_name, cache):
        lowered = item_name.lower()
        for category_name, keywords in CATEGORY_KEYWORDS.items():
            if any(keyword in lowered for keyword in keywords):
                if category_name not in cache:
                    category, _ = Category.objects.get_or_create(name=category_name)
                    cache[category_name] = category
                return cache[category_name]
        return None

    def _parse_quantity(self, value):
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return int(value)
        match = re.search(r"\d+", str(value))
        return int(match.group()) if match else None

    def _load_dimensions(self, wb):
        if "Box Count and Sizes" not in wb.sheetnames:
            return {}
        ws = wb["Box Count and Sizes"]
        dims = {}
        for box_number, size, dimensions, *_rest in ws.iter_rows(min_row=2, values_only=True):
            if box_number is None or not dimensions:
                continue
            try:
                dims[int(box_number)] = str(dimensions).strip()
            except (TypeError, ValueError):
                continue
        return dims
