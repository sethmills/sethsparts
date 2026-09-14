"""One-off export: parts needing review/clarification -> an editable spreadsheet.

Run with: ./venv/bin/python scripts/export_parts_review.py
Not a management command -- this is a one-time deliverable generator, not an
app feature.

The workbook it writes (`docs/parts_review.xlsx`) is a working file for the owner of this
workshop, not part of the app, so it is gitignored -- a clone of this repo gets the code,
not somebody else's inventory. Same for the clarification worklist beside it.
"""
import json
import os
import sys
from urllib.parse import quote

import django

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.worksheet import Worksheet  # noqa: E402

from inventory.models import Part, StockItem  # noqa: E402

SCRATCH = "/private/tmp/claude-501/-Users-seth/d95730c8-627a-4511-954d-f4112cc34ecf/scratchpad"
SOURCE_FILES = [
    "adafruit_parts.json",
    "batch1_results.json",
    "batch2_merged.json",
    "batch3_results.json",
]
SITE = "https://sethsparts.com"

# If set (via --merge <path>), carry over whatever Seth already typed into the
# two "fill in" columns of an earlier export, keyed by Part ID, so re-running
# this script doesn't clobber his in-progress work.
PRIOR_FILE = sys.argv[2] if len(sys.argv) > 2 and sys.argv[1] == "--merge" else None


def load_prior_answers(path):
    if not path or not os.path.exists(path):
        return {}
    import openpyxl as _openpyxl

    prior = {}
    wb = _openpyxl.load_workbook(path, data_only=True)
    for sheet_name in ("Needs Review", "Needs Clarification"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        # the two rightmost "(fill in)" columns, whatever their exact position
        fill_in_cols = [i for i, h in enumerate(header, start=1) if h and "fill in" in str(h)]
        id_col = header.index("ID") + 1
        for row in ws.iter_rows(min_row=2):
            part_id = row[id_col - 1].value
            if part_id is None:
                continue
            values = tuple(row[c - 1].value for c in fill_in_cols)
            if any(v not in (None, "") for v in values):
                prior[(sheet_name, int(part_id))] = values
    return prior


def locations_summary(part):
    locs = []
    for si in StockItem.objects.filter(part=part).select_related("container", "drawer").order_by(
        "container__number", "drawer__label"
    ):
        where = f"#{si.container.number}/{si.drawer.label}" if si.drawer else f"#{si.container.number}"
        locs.append(where)
    return "; ".join(locs)

HEADER_FILL = PatternFill("solid", fgColor="1E2128")
HEADER_FONT = Font(color="E6E8EB", bold=True)
LOW_FILL = PatternFill("solid", fgColor="FFF3CD")


def load_notes():
    notes = {}
    for fn in SOURCE_FILES:
        path = os.path.join(SCRATCH, fn)
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                for entry in json.load(f):
                    notes[entry["id"]] = entry
    return notes


def style_header(ws: Worksheet, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28


def search_url(name):
    return f"https://www.google.com/search?q={quote(name)}"


def part_url(pk):
    return f"{SITE}/parts/{pk}/"


def add_hyperlink(ws, row, col, url, label):
    cell = ws.cell(row=row, column=col, value=label)
    cell.hyperlink = url
    cell.font = Font(color="1F5FD2", underline="single")
    return cell


def build_needs_review_sheet(wb, notes, prior):
    ws = wb.active
    ws.title = "Needs Review"
    headers = [
        "ID", "Part Name", "Current Guess", "Confidence", "Why Uncertain", "Location",
        "Part Page", "Quick Search", "Correct Product / Model (fill in)", "Correct URL (fill in)",
    ]
    ws.append(headers)
    style_header(ws, len(headers))

    parts = Part.objects.filter(enrichment_status="needs_review").order_by("name")
    for i, p in enumerate(parts, start=2):
        entry = notes.get(p.id, {})
        ws.cell(row=i, column=1, value=p.id)
        ws.cell(row=i, column=2, value=p.name)
        ws.cell(row=i, column=3, value=entry.get("matched_product_name", p.description))
        confidence = entry.get("confidence", "")
        ws.cell(row=i, column=4, value=confidence)
        ws.cell(row=i, column=5, value=entry.get("notes", ""))
        ws.cell(row=i, column=6, value=locations_summary(p))
        add_hyperlink(ws, i, 7, part_url(p.id), "Open in app")
        add_hyperlink(ws, i, 8, search_url(p.name), "Google it")
        prior_values = prior.get(("Needs Review", p.id), ("", ""))
        ws.cell(row=i, column=9, value=prior_values[0])
        ws.cell(row=i, column=10, value=prior_values[1] if len(prior_values) > 1 else "")

        if confidence == "low":
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = LOW_FILL

    widths = [6, 32, 40, 11, 50, 20, 12, 12, 30, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{ws.max_row}"
    return ws


def build_needs_clarification_sheet(wb, prior):
    ws = wb.create_sheet("Needs Clarification")
    headers = [
        "ID", "Part Name", "Container", "Drawer", "Qty",
        "Part Page", "Quick Search", "What is this? (fill in)", "Correct URL (fill in)",
    ]
    ws.append(headers)
    style_header(ws, len(headers))

    parts = Part.objects.filter(enrichment_status="needs_clarification").order_by(
        "stock_items__container__number", "stock_items__drawer__label", "name"
    )
    seen = set()
    row = 2
    for p in parts:
        if p.id in seen:
            continue
        seen.add(p.id)
        prior_values = prior.get(("Needs Clarification", p.id), ("", ""))
        for si in StockItem.objects.filter(part=p).select_related("container", "drawer").order_by(
            "container__number", "drawer__label"
        ):
            ws.cell(row=row, column=1, value=p.id)
            ws.cell(row=row, column=2, value=p.name)
            ws.cell(row=row, column=3, value=si.container.number)
            ws.cell(row=row, column=4, value=si.drawer.label if si.drawer else "")
            ws.cell(row=row, column=5, value=si.quantity_raw or si.quantity)
            add_hyperlink(ws, row, 6, part_url(p.id), "Open in app")
            add_hyperlink(ws, row, 7, search_url(p.name), "Google it")
            ws.cell(row=row, column=8, value=prior_values[0])
            ws.cell(row=row, column=9, value=prior_values[1] if len(prior_values) > 1 else "")
            row += 1

    widths = [6, 32, 11, 11, 8, 12, 12, 34, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{ws.max_row}"
    return ws


def main():
    notes = load_notes()
    prior = load_prior_answers(PRIOR_FILE)
    if PRIOR_FILE:
        print(f"Merged {len(prior)} previously-answered rows from {PRIOR_FILE}")
    wb = Workbook()
    build_needs_review_sheet(wb, notes, prior)
    build_needs_clarification_sheet(wb, prior)
    out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "parts_review.xlsx")
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
